from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Stochastic_Pirates_Methodology_Submission.docx")

NAVY = "17365D"
LIGHT_BLUE = "EAF1F8"
PALE = "F7F9FB"
GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, white=False, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(255, 255, 255) if white else BLACK
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_margin(cell)
    borders(cell)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_text(doc, text, *, style=None, bold_lead=None, before=0, after=7):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r = p.add_run(text[len(bold_lead):])
    else:
        r = p.add_run(text)
    for run in p.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(10.5)
        run.font.color.rgb = BLACK
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(10.3)
    r.font.color.rgb = BLACK
    return p


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_cells = table.rows[0].cells
    for i, value in enumerate(headers):
        shade(header_cells[i], NAVY)
        set_cell_text(header_cells[i], value, bold=True, white=True, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_width(header_cells[i], widths[i])
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_pr.append(repeat_header)
    for row_no, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if row_no % 2 == 1:
                shade(cells[i], PALE)
            set_cell_text(cells[i], value, size=font_size)
            set_width(cells[i], widths[i])
    for row in table.rows:
        row._tr.get_or_add_trPr()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(5)

    for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Stochastic Pirates  |  Business Entity Resolution Methodology")
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(89, 89, 89)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    configure_styles(doc)
    add_footer(section)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Business Entity Resolution Approach")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    subtitle.paragraph_format.line_spacing = 1.05
    r = subtitle.add_run("Methodology submission for ML Challenge 2026")
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(89, 89, 89)

    add_table(doc,
        ["Submission details", "Value"],
        [("Team", "Stochastic Pirates"), ("Institution", "Pandit Deendayal Energy University")],
        [1.8, 5.1], 9.5)

    h = doc.add_paragraph("1. Executive Summary", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "Our solution is a local-only, supervised entity-resolution pipeline designed for a precision-sensitive macro F0.5 objective. For each Source-1 business record, the system retrieves a small, high-recall set of plausible Source-2 and Source-3 candidates, scores only those pairs with a gradient-boosted classifier, and applies source-specific decision thresholds selected on held-out Source-1 entities. The approach combines robust multilingual text normalization, complementary blocking signals, pairwise similarity features, hard-negative learning, and safeguards against unsupported matching assumptions.")
    add_text(doc, "The system uses only the challenge files. It does not query external databases, geocoding services, websites, APIs, or third-party business directories. This keeps the workflow reproducible, contest-compliant, and deployable in an offline environment.")

    h = doc.add_paragraph("2. End to End Methodology", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "The workflow separates retrieval from decision-making. Blocking keeps comparison tractable and recall-oriented; the learned scorer resolves ambiguity among the retrieved candidates; threshold tuning reflects the competition metric rather than an arbitrary probability cutoff.")
    add_table(doc,
        ["Stage", "Purpose", "Implementation"],
        [
            ("Normalize", "Create comparable representations of names and addresses.", "Unicode transliteration, lowercase normalization, punctuation and URL cleanup, legal-suffix removal, address-abbreviation expansion, and numeric-address normalization."),
            ("Block", "Retrieve plausible matches without all-pairs comparison.", "Country-partitioned inverted indices over name, address, number, and name-address composite keys; candidates are IDF-ranked and capped at the top 50."),
            ("Score", "Estimate whether a candidate pair represents the same business.", "One LightGBM binary classifier using 26 deterministic pair and candidate-context features."),
            ("Decide", "Convert calibrated scores into the final links.", "Separate Source-2 and Source-3 thresholds chosen by validation macro F0.5; optional one-to-one target assignment only when verified by the labels."),
        ], [1.0, 1.9, 4.0], 9.0)

    h = doc.add_paragraph("3. Candidate Generation and Blocking Strategy", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "Candidate generation is intentionally redundant: a genuine match can survive via a business-name key, an address-based key, or a name-address composite even when another field is incomplete or noisy. Records are compared only within the observed country field. This is an open-set partition rather than a hard-coded list of countries, so it naturally handles France and any additional country labels present in the data.")
    add_text(doc, "For each Source-2 and Source-3 target, we construct an inverted index from the following keys:", after=4)
    add_bullet(doc, "Canonical compact business name, legal-suffix-stripped name core, first two name tokens, and sorted name tokens.")
    add_bullet(doc, "Significant individual name tokens, with very common postings suppressed only for weak single-token key types.")
    add_bullet(doc, "Address number combined with rare address tokens or the final address token.")
    add_bullet(doc, "First business-name token combined with an address number or location token.")
    add_bullet(doc, "Pairs of rare address tokens.")
    add_text(doc, "Candidates are ranked by the sum of inverse-frequency weights of their shared keys. Therefore, agreement on several rare, independent signals receives more weight than agreement on one ubiquitous token. Broad postings are filtered only for noisy key families; exact and composite keys remain available. At inference time, the top 50 candidates per Source-1 record are sent to the scorer, and the submitted candidate-pair file is generated from that exact scored candidate list.")

    h = doc.add_paragraph("4. Model Architecture and Feature Engineering", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "The scoring model is a single LightGBM binary classifier with 650 trees, learning rate 0.035, 63 leaves, row and feature subsampling of 0.85, L2 regularization of 1.0, and a fixed random seed of 2026. It learns from labeled true pairs and candidate-derived hard negatives, rather than from randomly sampled non-matches that would be unrealistically easy to distinguish.")
    add_table(doc,
        ["Feature family", "Signals"],
        [
            ("Name similarity", "Exact normalized-name and core-name agreement; normalized Levenshtein; Jaro-Winkler; token-sort and token-set ratios; token Jaccard and overlap; length ratio; character-trigram Jaccard; first-token similarity."),
            ("Address similarity", "Address presence; exact, edit, Jaro-Winkler, token-sort, token-set, and token-Jaccard agreement after abbreviation expansion and cleanup."),
            ("Numeric structure", "Shared address-number indicator; number conflict; primary-number agreement and conflict. Leading zeros are normalized while the first observed number is retained as a separate anchor."),
            ("Interactions and context", "Name-by-address interaction; exact core name with a missing target address; IDF-weighted blocking score; and a Source-2 indicator."),
        ], [1.5, 5.4], 9.1)
    add_text(doc, "These features deliberately combine tolerant string similarity with contradiction signals. High name similarity alone is not sufficient when both records contain conflicting address numbers; conversely, a clean name and address agreement can compensate for superficial formatting variation.")

    h = doc.add_paragraph("5. Training Validation and Decision Policy", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "Validation is performed using a deterministic Source-1-level split. Each Source-1 ID is assigned using a fixed BLAKE2b hash; 15 percent are held out for threshold selection, while the remaining records provide training pairs. The implementation caps the training sample at 60,000 Source-1 entities and the validation sample at 10,000 entities so the procedure remains practical and repeatable. All candidate retrieval, scoring, and threshold selection for the validation set are performed after model fitting.")
    add_text(doc, "Separate Source-2 and Source-3 probability thresholds are selected by grid search from 0.40 to 0.95 in steps of 0.05, maximizing macro F0.5 on the held-out records. This makes false-positive links costly, which is appropriate for a precision-weighted challenge metric and for singleton records that should receive no match.")
    add_table(doc,
        ["Control", "Rationale"],
        [
            ("Hard-negative training", "Teaches the model to distinguish genuine matches from the most plausible false look-alikes returned by the blocker."),
            ("Candidate-recall measurement", "The training and validation procedures record the fraction of true links captured by the blocker, isolating retrieval performance from model performance."),
            ("Source-specific thresholds", "Allows Source-2 and Source-3 score distributions to be calibrated separately instead of enforcing one global cutoff."),
            ("Conditional target exclusivity", "Greedy highest-probability one-to-one assignment is used only when the supplied ground truth shows that a target ID is not linked to multiple Source-1 entities. Otherwise, all threshold-qualified links are retained."),
        ], [1.65, 5.25], 9.1)

    h = doc.add_paragraph("6. Reproducibility and Submission Integrity", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "The pipeline is deterministic under the documented seed and writes its trained model, feature metadata, selected thresholds, candidate-recall diagnostics, and validation metrics to the work directory. The final outputs are regenerated by the same run: matching_results.tsv and candidate_pairs.tsv. Candidate pairs are not an auxiliary approximation; they are the actual shortlist consumed by the classifier, so every reported match necessarily belongs to the submitted candidate set.")
    add_text(doc, "Dependencies are pinned: LightGBM 4.5.0, NumPy 1.26.4, RapidFuzz 3.9.7, scikit-learn 1.5.2, SciPy 1.14.1, and text-unidecode 1.3. The code includes a prediction-only mode that reloads the saved LightGBM model and metadata, which supports repeatable output generation after training.")
    add_text(doc, "Before submission, both TSV files should be regenerated from the final run and checked with the organizer-provided validation utility. This check confirms file structure and alignment before upload; the final archive should be built only after the validator reports success.")

    h = doc.add_paragraph("7. Additional Relevant Information", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "The system was designed to make conservative, explainable decisions. Its blockers use human-interpretable agreements such as a shared business-name core, rare address tokens, and address numbers. Its model features expose both supporting evidence and contradictions. Most importantly, it does not assume that target records are globally unique unless that property is demonstrated in the supplied ground truth. This avoids turning an undocumented dataset assumption into a source of systematic false negatives or false positives.")
    add_text(doc, "The implementation supports batching and writes intermediate accepted scores to a temporary SQLite database during prediction. This keeps memory usage bounded for large test sets while preserving deterministic processing and final input order in the submission files.")

    h = doc.add_paragraph("Appendix  Reproduction Command", style="Heading 1")
    keep_with_next(h)
    add_text(doc, "From the code/business_entity_resolution directory, place the supplied data in the expected dataset folders and run:", after=3)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run("python src/pipeline.py run --train-dir ..\\..\\..\\dataset\\train --test-dir ..\\..\\..\\dataset\\test --output-dir ..\\..\\..\\output --work-dir ..\\..\\..\\work")
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(8.5)
    r.font.color.rgb = BLACK
    add_text(doc, "The run performs validation, trains the model, selects thresholds, writes model artifacts and metrics, and creates the two required TSV files.")

    core = doc.core_properties
    core.title = "Business Entity Resolution Approach"
    core.subject = "Methodology submission"
    core.author = "Stochastic Pirates"
    core.comments = ""
    doc.save(OUT)


if __name__ == "__main__":
    build()
